"""Automatic workbook creation follows successful stack completion only."""

from pathlib import Path
import queue
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest import mock

import numpy as np
from openpyxl import Workbook, load_workbook

from aidas.steps import step4_analyze_isez as module
from aidas.steps.step5_compile import DEFAULT_OUTPUT_FILENAME, Step5Frame
from tests.steps.step4 import test_step4_batch_stack_building as batch_fixtures


class _Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class Step4AutomaticExcelTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name).resolve()
        self.info = self.enterContext(mock.patch.object(module.messagebox, "showinfo"))
        self.error = self.enterContext(mock.patch.object(module.messagebox, "showerror"))

    def _single_frame(self):
        frame = object.__new__(module.Step4Frame)
        frame.status_var = _Var()
        frame.batch_roi_notebook = None
        frame.batch_roi_tab_states = {}
        frame._active_batch_roi_tab = None
        frame.batch_roi_paths = []
        frame.batch_roi_index = -1
        frame.batch_roi_root = None
        frame._stack_building = False
        frame._stack_build_complete = False
        frame.output_dir_var = _Var(str(self.root))
        frame.input_dir_var = _Var(str(self.root))
        frame.image = np.ones((8, 12), dtype=np.int16)
        frame.rois = [SimpleNamespace(suffix="01")]
        frame.completed = {"01": SimpleNamespace(roi=frame.rois[0])}
        frame.on_continue_to_step5 = mock.Mock()
        for name in (
            "_show_processing_complete", "_show_processed_grid_notice",
            "_show_stack_building_notice", "_hide_grid_notice",
            "_update_build_stack_button_state", "_update_auto_detect_button_state",
            "_update_continue_to_step5_button_state", "_write_stack_output_files",
        ):
            setattr(frame, name, mock.Mock())
        frame._workers = []
        frame._callbacks = []
        frame._start_background_worker = frame._workers.append
        frame.after = lambda _delay, callback: frame._callbacks.append(callback) or "poll"
        return frame

    def _batch_fixture(self, *, all_ready=True):
        fixture = batch_fixtures.BatchStackBuildingTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        frame = fixture.frame
        frame.batch_roi_root = fixture.root
        frame.batch_roi_paths = []
        frame.on_continue_to_step5 = mock.Mock()
        if all_ready:
            fixture.states["unfinished"]["completed"]["02"] = SimpleNamespace(roi=frame.rois[1])
        return fixture

    @staticmethod
    def _finish_single(frame):
        frame._workers.pop(0)()
        frame._callbacks.pop(0)()

    def test_single_stack_starts_excel_after_output_worker_finishes_only_once(self):
        frame = self._single_frame()
        frame._build_stack_outputs()
        frame.on_continue_to_step5.assert_not_called()
        poll = frame._callbacks[0]

        self._finish_single(frame)
        poll()  # A queued callback for the finished run must be harmless.

        frame._write_stack_output_files.assert_called_once()
        self.assertFalse(frame._stack_building)
        self.assertTrue(frame._stack_build_complete)
        frame.on_continue_to_step5.assert_called_once_with(str(self.root))
        self.info.assert_not_called()

    def test_all_ready_batch_tabs_finish_before_one_excel_handoff(self):
        fixture = self._batch_fixture()
        fixture.frame._build_stack_outputs()
        fixture.frame.on_continue_to_step5.assert_not_called()
        poll = fixture.callbacks[0]

        fixture.finish()
        poll()

        self.assertTrue(all(state["complete"] for state in fixture.states.values()))
        self.assertEqual(fixture.frame._write_stack_output_files.call_count, 3)
        fixture.frame.on_continue_to_step5.assert_called_once_with(fixture.root)

    def test_unfinished_batch_tab_prevents_excel_after_ready_tabs_finish(self):
        fixture = self._batch_fixture(all_ready=False)
        fixture.frame._build_stack_outputs()
        fixture.finish()

        self.assertTrue(fixture.states["active"]["complete"])
        self.assertFalse(fixture.states["unfinished"]["complete"])
        fixture.frame.on_continue_to_step5.assert_not_called()

    def test_partially_failed_batch_does_not_generate_excel(self):
        fixture = self._batch_fixture()
        fixture.frame._write_stack_output_files.side_effect = [OSError("Disk full"), None, None]
        fixture.frame._build_stack_outputs()
        fixture.finish()

        self.assertFalse(fixture.states["active"]["complete"])
        self.assertTrue(fixture.states["ready"]["complete"])
        fixture.error.assert_called_once()
        fixture.frame.on_continue_to_step5.assert_not_called()

    def test_cancelled_batch_does_not_write_or_generate_excel(self):
        fixture = self._batch_fixture()
        fixture.frame._build_stack_outputs()
        fixture.frame._cancel_batch_stack_build()
        fixture.finish()

        fixture.frame._write_stack_output_files.assert_not_called()
        fixture.frame.on_continue_to_step5.assert_not_called()

    def test_changed_roi_snapshot_prevents_stale_batch_handoff(self):
        fixture = self._batch_fixture()
        fixture.frame._build_stack_outputs()
        fixture.workers.pop()()
        fixture.states["active"]["completed"] = {}
        fixture.callbacks.pop(0)()

        self.assertFalse(fixture.states["active"]["complete"])
        fixture.frame.on_continue_to_step5.assert_not_called()

    def test_detached_single_worker_cannot_trigger_excel(self):
        frame = self._single_frame()
        frame._build_stack_outputs()
        frame._stack_build_events = None
        self._finish_single(frame)

        frame.on_continue_to_step5.assert_not_called()
        self.assertFalse(frame._stack_build_complete)

    def test_failed_single_stack_does_not_trigger_excel(self):
        frame = self._single_frame()
        frame._write_stack_output_files.side_effect = OSError("Read-only output")
        frame._build_stack_outputs()
        self._finish_single(frame)

        self.error.assert_called_once()
        frame.on_continue_to_step5.assert_not_called()
        self.assertFalse(frame._stack_build_complete)

    def test_legacy_batch_hands_off_after_last_folder_only_once(self):
        frame = self._single_frame()
        frame.batch_roi_root = self.root
        frame.batch_roi_paths = [self.root / "first" / "input.hdr", self.root / "last" / "input.hdr"]
        frame.batch_roi_index = 0
        frame.batch_roi_skipped = 0
        frame._load_path = mock.Mock(side_effect=lambda _path: setattr(frame, "_stack_build_complete", False))

        frame._finish_stack_build(self.root / "first")
        frame.on_continue_to_step5.assert_not_called()
        frame._load_path.assert_called_once_with(frame.batch_roi_paths[1])
        frame._finish_stack_build(self.root / "last")
        frame._load_next_batch_roi()

        frame.on_continue_to_step5.assert_called_once_with(self.root)

    def test_next_single_build_can_generate_a_fresh_workbook(self):
        frame = self._single_frame()
        for generation in range(1, 3):
            frame._build_stack_outputs()
            self._finish_single(frame)
            self.assertEqual(frame._step5_handoff_generation, generation)

        self.assertEqual(frame.on_continue_to_step5.call_count, 2)

    def test_next_batch_build_rearms_handoff_after_previous_completion(self):
        fixture = self._batch_fixture()
        fixture.frame._build_stack_outputs()
        fixture.finish()
        self.assertEqual(fixture.frame._step5_handoff_generation, 1)
        fixture.states["active"]["complete"] = False
        fixture.frame._build_stack_outputs()
        fixture.finish()
        self.assertEqual(fixture.frame._step5_handoff_generation, 2)

        self.assertEqual(fixture.frame.on_continue_to_step5.call_count, 2)

    def test_standalone_step4_completes_without_step5_callback(self):
        frame = self._single_frame()
        frame.on_continue_to_step5 = None
        frame._build_stack_outputs()
        self._finish_single(frame)

        self.assertTrue(frame._stack_build_complete)
        frame._show_processing_complete.assert_called_once()
        self.info.assert_not_called()

    def test_stack_completion_creates_real_excel_without_confirmation_or_save_dialog(self):
        frame = self._single_frame()
        side = self.root / "Control" / "Subject_A" / "nasal"
        side.mkdir(parents=True)
        frame.output_dir_var.set(str(side))

        def write_measurements(_image, _ordered, outdir):
            workbook = Workbook()
            try:
                workbook.active.cell(1, 7, "Round")
                workbook.active.cell(2, 7, 10)
                workbook.save(outdir / "rr_MCPAR.xlsx")
            finally:
                workbook.close()

        frame._write_stack_output_files.side_effect = write_measurements
        step5 = object.__new__(Step5Frame)
        step5._running = False
        step5.after = mock.Mock(return_value="compile-poll")
        step5._output_user_selected = False
        step5.input_var = _Var()
        step5.output_var = _Var()
        step5.include_fovea_var = _Var(True)
        step5.status_var = _Var()
        step5._events = queue.Queue()
        step5._clear_log = mock.Mock()
        step5._append_log = mock.Mock()
        step5._set_running = mock.Mock(side_effect=lambda state: setattr(step5, "_running", state))
        step5.on_compilation_complete = mock.Mock()
        step5.source_step = frame
        frame.on_continue_to_step5.side_effect = step5.start_from_step4
        with mock.patch("tkinter.filedialog.asksaveasfilename") as save_as, mock.patch(
            "tkinter.messagebox.askyesno",
        ) as confirmation:
            frame._build_stack_outputs()
            self._finish_single(frame)
            step5._worker.join(timeout=10)
            self.assertFalse(step5._worker.is_alive())
            step5._drain_compile_events()

        output = side.parent.parent / DEFAULT_OUTPUT_FILENAME
        save_as.assert_not_called()
        confirmation.assert_not_called()
        self.info.assert_not_called()
        self.error.assert_not_called()
        step5.on_compilation_complete.assert_called_once_with(output)
        self.assertIn("Excel file saved", step5.status_var.get())
        workbook = load_workbook(output, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["rrMCP-AR", "ELM-RPE", "ONL", "Extraction_log"])
            self.assertEqual(workbook["rrMCP-AR"]["A1"].value, "Control")
            self.assertEqual(workbook["rrMCP-AR"]["A3"].value, "Subject_A")
            self.assertEqual(workbook["rrMCP-AR"]["B3"].value, 10)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
