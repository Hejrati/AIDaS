from __future__ import annotations

import inspect
from pathlib import Path
import queue
from types import SimpleNamespace
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

from aidas.app import AIDaSApp
from aidas.steps.step5_compile import DEFAULT_OUTPUT_FILENAME, Step5Frame
from aidas.utils.ui_utils import _compile_icon_image


class _Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class Step5CompileTests(unittest.TestCase):
    def test_compile_icon_uses_the_vaadin_grid_shape(self):
        icon = _compile_icon_image("#FFFFFF", 64)

        self.assertEqual(icon.size, (64, 64))
        self.assertEqual(icon.getpixel((0, 0))[3], 0)
        self.assertEqual(icon.getpixel((4, 8)), (255, 255, 255, 255))
        self.assertEqual(icon.getpixel((54, 2)), (255, 255, 255, 255))
        self.assertEqual(icon.getpixel((47, 0))[3], 0)

    def test_step5_compile_button_uses_dedicated_compile_icon(self):
        source = inspect.getsource(Step5Frame._build_ui)

        self.assertIn("load_compile_ctk_icon", source)
        self.assertNotIn("flat-color-icons--process.png", source)

    def test_success_notifies_app_after_the_workbook_is_saved(self):
        frame = object.__new__(Step5Frame)
        frame.output_var = _Var("requested.xlsx")
        frame.status_var = _Var()
        frame._set_running = mock.Mock()
        frame._append_log = mock.Mock()
        frame.on_compilation_complete = mock.Mock()
        result = SimpleNamespace(output_path=Path("compiled.xlsx"))

        with mock.patch("aidas.steps.step5_compile.messagebox.showinfo") as popup:
            frame._compilation_succeeded(result)

        popup.assert_not_called()
        frame.on_compilation_complete.assert_called_once_with(Path("compiled.xlsx"))

    def test_final_success_updates_status_without_a_completion_popup(self):
        app = AIDaSApp.__new__(AIDaSApp)
        app._set_status_message = mock.Mock()
        app._on_step5_compilation_complete(Path("compiled.xlsx"))
        app._set_status_message.assert_called_once_with("Excel file created successfully.")
        self.assertNotIn("_workflow_completion_popup", app.__dict__)

    def test_handoff_folder_prefills_input_and_default_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            frame = object.__new__(Step5Frame)
            frame.input_var = _Var()
            frame.output_var = _Var()
            frame.status_var = _Var()
            frame._output_user_selected = False

            accepted = frame.set_input_folder(temp_dir)

            self.assertTrue(accepted)
            folder = Path(temp_dir).resolve()
            self.assertEqual(Path(frame.input_var.get()), folder)
            self.assertEqual(
                Path(frame.output_var.get()),
                folder / DEFAULT_OUTPUT_FILENAME,
            )
            self.assertIn("Ready to create an Excel file", frame.status_var.get())

    def test_validation_adds_xlsx_extension(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            frame = object.__new__(Step5Frame)
            frame.input_var = _Var(temp_dir)
            frame.output_var = _Var(str(Path(temp_dir) / "combined"))

            paths = frame._validated_paths()

            folder = Path(temp_dir).resolve()
            self.assertEqual(paths, (folder, folder / "combined.xlsx"))
            self.assertEqual(Path(frame.output_var.get()).suffix, ".xlsx")

    def test_worker_uses_existing_compiler_service(self):
        frame = object.__new__(Step5Frame)
        frame._events = queue.Queue()
        root = Path("root")
        output = Path("compiled.xlsx")
        result = SimpleNamespace(output_path=output)

        with mock.patch(
            "aidas.services.step4_compiler.compile_step4_results",
            return_value=result,
        ) as compile_results:
            frame._compile_worker(root, output, True)

        compile_results.assert_called_once_with(
            root,
            output,
            include_fovea=True,
            log_callback=mock.ANY,
            progress_callback=mock.ANY,
        )
        self.assertEqual(frame._events.get_nowait(), ("success", result))

    def test_app_handoff_selects_step5_and_starts_compilation(self):
        app = AIDaSApp.__new__(AIDaSApp)
        step5 = mock.Mock()
        app.__dict__["step5"] = step5
        app.notebook = mock.Mock()
        app.update_idletasks = mock.Mock()
        folder = Path("measurement-root")

        app._on_step4_continue_to_step5(folder)

        app.notebook.select.assert_called_once_with(step5)
        app.update_idletasks.assert_called_once_with()
        step5.start_from_step4.assert_called_once_with(folder)

    def test_switching_directly_from_step4_to_step5_does_not_transfer_or_compile(self):
        app = AIDaSApp.__new__(AIDaSApp)
        app.step5 = mock.Mock()
        app.notebook = mock.Mock()
        app.notebook.index.return_value = 4
        app._last_workflow_tab_index = 3
        app._refresh_image_path_status = mock.Mock()
        app._sync_workflow_progress_strip = mock.Mock()

        app._on_workflow_tab_changed()
        app._on_workflow_tab_changed()

        self.assertEqual(app.step5.on_show.call_count, 2)
        app.step5.start_from_step4.assert_not_called()

    def test_showing_step5_preserves_its_input_and_never_reads_step4_folders(self):
        for folder in ("", "previously-selected-folder"):
            with self.subTest(folder=folder):
                frame = object.__new__(Step5Frame)
                frame.input_var = _Var(folder)
                frame._source_input_folder = mock.Mock()
                frame.set_input_folder = mock.Mock()
                frame._run_clicked = mock.Mock()
                frame._refresh_directory_view = mock.Mock()
                frame.on_show()
                self.assertEqual(frame.input_var.get(), folder)
                frame._source_input_folder.assert_not_called()
                frame.set_input_folder.assert_not_called()
                frame._run_clicked.assert_not_called()
                frame._refresh_directory_view.assert_called_once_with()

    def test_opening_step5_from_another_tab_does_not_start_compilation(self):
        app = AIDaSApp.__new__(AIDaSApp)
        app.step5 = mock.Mock()
        app.notebook = mock.Mock()
        app.notebook.index.return_value = 4
        app._last_workflow_tab_index = 0
        app._refresh_image_path_status = mock.Mock()
        app._sync_workflow_progress_strip = mock.Mock()

        app._on_workflow_tab_changed()

        app.step5.on_show.assert_called_once_with()
        app.step5.start_from_step4.assert_not_called()

    def test_explicit_handoff_and_queued_tab_event_start_only_once(self):
        app = AIDaSApp.__new__(AIDaSApp)
        app.step5 = mock.Mock()
        app.notebook = mock.Mock()
        app.notebook.index.return_value = 4
        app._last_workflow_tab_index = 3
        app._refresh_image_path_status = mock.Mock()
        app._sync_workflow_progress_strip = mock.Mock()
        app.update_idletasks = mock.Mock(side_effect=app._on_workflow_tab_changed)
        folder = Path("measurement-root")

        app._on_step4_continue_to_step5(folder)
        app._on_workflow_tab_changed()

        app.step5.start_from_step4.assert_called_once_with(folder)

    def _handoff_frame(self, root):
        frame = object.__new__(Step5Frame)
        frame._running = False
        frame._output_user_selected = False
        frame.after = mock.Mock(return_value="compile-poll")
        frame.input_var = _Var("old-input")
        frame.output_var = _Var()
        frame.status_var = _Var()
        frame.include_fovea_var = _Var(True)
        frame._start_compiler = mock.Mock()
        frame.source_step = SimpleNamespace(batch_roi_root=root)
        return frame

    def test_handoff_automatically_runs_with_the_new_input_and_default_output(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            frame = self._handoff_frame(root)
            frame.start_from_step4(root)
            frame._start_compiler.assert_called_once_with(
                root, root / DEFAULT_OUTPUT_FILENAME, True,
                thread_name="aidas-step5-compiler",
            )

    def test_handoff_promotes_side_and_subject_folders_to_the_group(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            group = Path(temp_dir).resolve() / "Control"
            subject = group / "Subject_A"
            side = subject / "nasal"
            nested = side / "exports"
            nested.mkdir(parents=True)
            for selected in (side, subject, nested):
                with self.subTest(selected=selected):
                    frame = self._handoff_frame(None)
                    frame.start_from_step4(selected)
                    frame._start_compiler.assert_called_once_with(
                        group, group / DEFAULT_OUTPUT_FILENAME, True,
                        thread_name="aidas-step5-compiler",
                    )

    def test_handoff_infers_common_root_of_all_step4_tabs(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            paths = [
                root / "Control" / "Subject_A" / "nasal" / "_flat_LIGHT.hdr",
                root / "Treatment" / "Subject_B" / "temporal" / "_flat_LIGHT.hdr",
            ]
            for path in paths:
                path.parent.mkdir(parents=True)
            frame = self._handoff_frame(None)
            frame.source_step.batch_roi_paths = paths
            frame.start_from_step4()
            frame._start_compiler.assert_called_once_with(
                root, root / DEFAULT_OUTPUT_FILENAME, True,
                thread_name="aidas-step5-compiler",
            )

    def test_handoff_preserves_custom_output_and_fovea_choice(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            frame = self._handoff_frame(root)
            output = root / "custom.xlsx"
            frame.output_var.set(str(output))
            frame._output_user_selected = True
            frame.include_fovea_var.set(False)
            frame.start_from_step4()
            frame._start_compiler.assert_called_once_with(
                root, output, False, thread_name="aidas-step5-compiler",
            )

    def test_handoff_does_not_start_a_second_worker_while_running(self):
        frame = self._handoff_frame(None)
        frame._running = True
        frame.set_input_folder = mock.Mock()
        frame.start_from_step4("new-root")
        frame.set_input_folder.assert_not_called()
        frame._start_compiler.assert_not_called()

    def test_completed_step4_runs_are_queued_without_overlapping_workers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            first = root / "first"
            latest = root / "latest"
            first.mkdir()
            latest.mkdir()
            frame = self._handoff_frame(None)
            frame._running = True
            frame._drain_compile_events = mock.Mock()
            frame.winfo_exists = mock.Mock(return_value=True)
            frame.after = mock.Mock()
            frame._start_compiler.side_effect = lambda *_args, **_kwargs: setattr(frame, "_running", True)

            frame.start_from_step4(first)
            frame.start_from_step4(latest)
            frame.start_from_step4(latest)
            frame._poll_events()
            frame._start_compiler.assert_not_called()

            frame._running = False
            frame._poll_events()
            frame._poll_events()
            frame._start_compiler.assert_called_once_with(
                first, first / DEFAULT_OUTPUT_FILENAME, True,
                thread_name="aidas-step5-compiler",
            )
            frame._running = False
            frame._poll_events()
            frame._running = False
            frame._poll_events()
            self.assertEqual(frame._start_compiler.call_args_list, [mock.call(
                first, first / DEFAULT_OUTPUT_FILENAME, True,
                thread_name="aidas-step5-compiler",
            ), mock.call(
                latest, latest / DEFAULT_OUTPUT_FILENAME, True,
                thread_name="aidas-step5-compiler",
            )])

    def test_a_queued_export_does_not_start_after_step5_is_destroyed(self):
        frame = self._handoff_frame(None)
        frame._running = True
        frame.start_from_step4("latest-input")
        frame._running = False
        frame._drain_compile_events = mock.Mock()
        frame.winfo_exists = mock.Mock(return_value=False)
        frame.after = mock.Mock()

        frame._poll_events()

        frame._start_compiler.assert_not_called()

    def test_visiting_step5_before_any_step4_file_is_loaded_does_not_compile_defaults(self):
        frame = self._handoff_frame(None)
        frame.source_step.output_dir_var = _Var("default-output-directory")
        self.assertEqual(frame._source_input_folder(), "")
        frame.start_from_step4()
        frame._start_compiler.assert_not_called()

    def test_invalid_handoff_does_not_compile_the_previous_input(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            frame = self._handoff_frame(root)
            frame.input_var.set(str(root))
            frame.start_from_step4(root / "missing")
            frame._start_compiler.assert_not_called()
            self.assertIn("unavailable", frame.status_var.get())

    def test_missing_side_folder_does_not_expand_to_a_stale_existing_group(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            subject = root / "Control" / "Subject_A"
            subject.mkdir(parents=True)
            frame = self._handoff_frame(root)
            frame.start_from_step4(subject / "nasal")
            frame._start_compiler.assert_not_called()
            self.assertIn("unavailable", frame.status_var.get())

    def test_automatic_handoff_worker_creates_a_real_grouped_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir).resolve()
            for group, value in (("Control", 10), ("Treatment", 20)):
                side = root / group / "Subject_A" / "nasal"
                side.mkdir(parents=True)
                workbook = Workbook()
                workbook.active.cell(1, 7, "Round")
                workbook.active.cell(2, 7, value)
                workbook.save(side / "rr_MCPAR.xlsx")
                workbook.close()

            frame = self._handoff_frame(root)
            del frame._start_compiler  # Exercise the actual background worker.
            frame._events = queue.Queue()
            frame._clear_log = mock.Mock()
            frame._append_log = mock.Mock()
            frame._set_running = mock.Mock(side_effect=lambda state: setattr(frame, "_running", state))
            frame.on_compilation_complete = mock.Mock()

            with mock.patch("aidas.steps.step5_compile.messagebox") as dialogs, mock.patch(
                "aidas.steps.step5_compile.filedialog",
            ) as file_dialogs:
                frame.start_from_step4()
                frame._worker.join(timeout=10)
                self.assertFalse(frame._worker.is_alive())
                frame._drain_compile_events()

            output = root / DEFAULT_OUTPUT_FILENAME
            self.assertEqual(dialogs.mock_calls, [])
            self.assertEqual(file_dialogs.mock_calls, [])
            self.assertFalse(frame._running)
            frame.on_compilation_complete.assert_called_once_with(output)
            self.assertIn("Excel file saved", frame.status_var.get())
            workbook = load_workbook(output, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["rrMCP-AR", "ELM-RPE", "ONL", "Extraction_log"])
                sheet = workbook["rrMCP-AR"]
                self.assertEqual([sheet[cell].value for cell in ("A1", "A3", "B3")], ["Control", "Subject_A", 10])
                self.assertEqual([sheet[cell].value for cell in ("E1", "E3", "F3")], ["Treatment", "Subject_A", 20])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
