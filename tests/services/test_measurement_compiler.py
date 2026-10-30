from __future__ import annotations

import inspect
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

from aidas.services.step4_compiler import (
    DEFAULT_OUTPUT_FILENAME,
    MeasureCompilation,
    clean_light_thickness_file,
    compile_step4_results,
    extract_elmrpe_onl_from_subject,
    find_light_files,
    find_subject_folders,
    find_results_file,
    find_thickness_file,
    make_light2_from_light,
    read_tsv_rows,
)


def _write_results(path: Path, first_value: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "in"
    worksheet.cell(1, 7).value = "Round"
    for offset in range(20):
        worksheet.cell(offset + 2, 7).value = first_value + offset
    workbook.save(path)
    workbook.close()


def _write_thickness(path: Path, elm_value: float, onl_value: float) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.write("Distance\tWhole\tELM\tONL\n")
        for distance in range(-1150, 1150):
            handle.write(f"{distance}\t0\t{elm_value}\t{onl_value}\n")


def _write_fovea(path: Path, value: float) -> None:
    path.write_text(f"foveaLIGHT 1 2\nRPEtoOLM {value} {value}\n", encoding="utf-8")


def _build_subject(root: Path) -> Path:
    subject = root / "LE" / "7_0_0"
    nasal = subject / "nasal"
    temporal = subject / "temporal"
    nasal.mkdir(parents=True)
    temporal.mkdir(parents=True)

    _write_results(nasal / "rr_MCPAR.xlsx", 1)
    _write_results(temporal / "Results.xlsx", 101)
    _write_thickness(nasal / "_thickness_vs_distance_from_fovea_LIGHT.txt", 10.0, 20.0)
    _write_thickness(temporal / "_thickness_vs_distance_from_fovea_LIGHT.txt", 30.0, 40.0)
    _write_fovea(nasal / "_fovea_light_profiles_LIGHT.txt", 55.77)
    _write_fovea(temporal / "_fovea_light_profiles_LIGHT.txt", 66.88)
    return subject


class Step4CompilerTests(unittest.TestCase):
    def test_compiler_uses_public_sheet_reordering_api(self):
        source = inspect.getsource(compile_step4_results)

        self.assertIn("workbook.move_sheet", source)
        self.assertNotIn("workbook._sheets", source)

    def test_full_numeric_subject_names_are_preserved_and_sorted_naturally(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            for subject_id in ("1000", "999"):
                nasal = root / "LE" / f"{subject_id}_0_0" / "nasal"
                nasal.mkdir(parents=True)
                _write_results(nasal / "Results.xlsx", 1)

            folders = find_subject_folders(root)
            result = compile_step4_results(root, root / "compiled.xlsx")

            self.assertEqual(set(folders["LE"]), {"999_0_0", "1000_0_0"})
            self.assertEqual(result.group_ids, {"LE": ("999_0_0", "1000_0_0")})

    def test_light_cleanup_preserves_original_and_removes_temporary_file_on_os_error(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "LIGHT.txt"
            source.write_bytes(b"Distance\tWhole\tELM\tONL\r\n-1\t0\tNA\t9\r\n0\t0\t1\t2\r\n")
            original = source.read_bytes()

            for operation in ("write_tsv_rows", "os.replace"):
                with self.subTest(operation=operation), mock.patch(
                    "aidas.services.step4_compiler." + operation,
                    side_effect=OSError("disk full"),
                ):
                    with self.assertRaisesRegex(OSError, "disk full"):
                        clean_light_thickness_file(source)

                self.assertEqual(source.read_bytes(), original)
                self.assertEqual(list(Path(temp_dir).iterdir()), [source])

    def test_light_cleanup_uses_fixed_elm_column_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            source.write_text(
                "Distance\tWhole\tELM\tONL\n"
                "-3\t0\tNA\t999\n"
                "-2\t0\t20\tNA\n"
                "-1\t0\t30\t300\n"
                "0\t0\tNA\t400\n"
                "1\t0\tNA\t500\n"
                "0\t0\tNA\t600\n"
                "2\t0\t40\t700\n"
                "3\tshort\n",
                encoding="utf-8",
            )

            self.assertEqual(make_light2_from_light(source, 4), str(source))
            self.assertEqual(read_tsv_rows(source), [
                ["Distance", "Whole", "ELM", "ONL"],
                ["-2", "0", "20", "NA"],
                ["-1", "0", "30", "300"],
                ["0", "0", "NA", "400"],
                ["2", "0", "40", "700"],
            ])
            cleaned = source.read_bytes()
            modified = source.stat().st_mtime_ns
            with mock.patch("aidas.services.step4_compiler.write_tsv_rows") as write:
                self.assertEqual(clean_light_thickness_file(source), str(source))
            write.assert_not_called()
            self.assertEqual(source.read_bytes(), cleaned)
            self.assertEqual(source.stat().st_mtime_ns, modified)
            self.assertEqual(list(Path(temp_dir).iterdir()), [source])

    def test_already_clean_light_retains_exact_bytes_without_a_rewrite(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "LIGHT.txt"
            original = b"Distance\tWhole\tELM\tONL\r\n-1\t0\t1\t2\r\n0\t0\tNA\t3\r\n"
            source.write_bytes(original)

            with mock.patch("aidas.services.step4_compiler.tempfile.mkstemp") as create:
                self.assertEqual(clean_light_thickness_file(source), str(source))

            create.assert_not_called()
            self.assertEqual(source.read_bytes(), original)

    def test_current_light_is_preferred_over_stale_light2_in_direct_and_recursive_paths(self):
        cases = (
            ("nasal", "nasal"),
            ("nasal", "nasal/exports/final"),
            ("rrMCPAR", "nasal/exports/final"),
            ("nasal", "other exports"),
        )
        for legacy_directory, current_directory in cases:
            with self.subTest(legacy=legacy_directory, current=current_directory), \
                    tempfile.TemporaryDirectory() as temp_dir:
                subject = Path(temp_dir)
                legacy = subject / legacy_directory / "_thickness_vs_distance_from_fovea_LIGHT2.txt"
                current = subject / current_directory / "_thickness_vs_distance_from_fovea_LIGHT.txt"
                legacy.parent.mkdir(parents=True, exist_ok=True)
                current.parent.mkdir(parents=True, exist_ok=True)
                legacy.write_text("stale", encoding="utf-8")
                current.write_text("current", encoding="utf-8")

                self.assertEqual(Path(find_thickness_file(subject)), current)
                nasal, _temporal = find_light_files(subject)
                self.assertEqual(Path(nasal), current)

    def test_current_temporal_light_is_preferred_over_direct_legacy_light2(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            subject = Path(temp_dir)
            side = subject / "temporal"
            current = side / "exports" / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            current.parent.mkdir(parents=True)
            current.write_text("current", encoding="utf-8")
            (side / "_thickness_vs_distance_from_fovea_LIGHT2.txt").write_text("stale", encoding="utf-8")

            _nasal, temporal = find_light_files(subject)

            self.assertEqual(Path(temporal), current)

    def test_legacy_light2_remains_a_read_only_fallback(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            subject = Path(temp_dir)
            side = subject / "nasal"
            side.mkdir()
            legacy = side / "_thickness_vs_distance_from_fovea_LIGHT2.txt"
            original = b"Distance\tWhole\tELM\tONL\r\n-1\t0\tNA\t100\r\n0\t0\t20\t200\r\n"
            legacy.write_bytes(original)

            self.assertEqual(Path(find_thickness_file(subject)), legacy)
            self.assertEqual(clean_light_thickness_file(legacy), str(legacy))
            blocks, _temporal, filename, _other, count, _other_count = extract_elmrpe_onl_from_subject(
                subject, "ONL", False,
            )

            self.assertEqual(Path(filename), legacy)
            self.assertEqual(blocks[0], 150)
            self.assertEqual(count, 2)
            self.assertEqual(legacy.read_bytes(), original)
            self.assertEqual(list(side.iterdir()), [legacy])

    def test_legacy_thickness_is_preferred_to_unrelated_current_light_profiles(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            subject = Path(temp_dir)
            side = subject / "nasal"
            side.mkdir()
            legacy = side / "_thickness_vs_distance_from_fovea_LIGHT2.txt"
            original = b"Distance\tWhole\tELM\tONL\n-1\t0\t10\t20\n0\t0\t30\t40\n"
            legacy.write_bytes(original)
            profile = side / "_fovea_light_profiles_LIGHT.txt"
            _write_fovea(profile, 55.77)
            profile_bytes = profile.read_bytes()

            self.assertEqual(Path(find_thickness_file(side)), legacy)
            nasal, _temporal = find_light_files(subject)
            self.assertEqual(Path(nasal), legacy)
            blocks, _temporal, filename, _other, count, _other_count = extract_elmrpe_onl_from_subject(
                subject, "ONL", False,
            )

            self.assertEqual(Path(filename), legacy)
            self.assertEqual(blocks[0], 30)
            self.assertEqual(count, 2)
            self.assertEqual(legacy.read_bytes(), original)
            self.assertEqual(profile.read_bytes(), profile_bytes)

    def test_profile_only_folders_are_not_cleaned_or_used_as_thickness_tables(self):
        for filename in ("_light_profiles_LIGHT.txt", "_fovea_light_profiles_LIGHT.txt",
                         "_fovea_light_LIGHT.txt", "_fovea_light_profiles_LIGHT2.txt"):
            with self.subTest(filename=filename), tempfile.TemporaryDirectory() as temp_dir:
                subject = Path(temp_dir)
                side = subject / "nasal"
                side.mkdir()
                profile = side / filename
                original = b"foveaLIGHT\t1\t2\r\nRPEtoOLM\tNA\tNA\r\n"
                profile.write_bytes(original)

                self.assertIsNone(find_thickness_file(side))
                self.assertEqual(find_light_files(subject), (None, None))
                for measurement in ("ELM-RPE", "ONL"):
                    nasal, temporal, nasal_file, temporal_file, nasal_count, temporal_count = (
                        extract_elmrpe_onl_from_subject(subject, measurement, False)
                    )
                    self.assertTrue(all(value is None for value in nasal + temporal))
                    self.assertIsNone(nasal_file)
                    self.assertIsNone(temporal_file)
                    self.assertEqual((nasal_count, temporal_count), (0, 0))
                self.assertEqual(profile.read_bytes(), original)
                self.assertEqual(list(side.iterdir()), [profile])

    def test_mixed_missing_elm_and_onl_rows_keep_former_compilation_values_without_light2(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            side = root / "Control" / "Subject_1" / "nasal"
            side.mkdir(parents=True)
            _write_results(side / "rr_MCPAR.xlsx", 1)
            source = side / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            source.write_text(
                "Distance\tWhole\tELM\tONL\n"
                "-3\t0\tNA\t999\n"
                "-2\t0\t20\tNA\n"
                "-1\t0\t30\t300\n"
                "0\t0\tNA\t400\n"
                "1\t0\tNA\t500\n"
                "0\t0\tNA\t600\n"
                "2\t0\t40\t700\n",
                encoding="utf-8",
            )
            output = root / "compiled.xlsx"

            compile_step4_results(root, output, include_fovea=False)

            workbook = load_workbook(output, data_only=True)
            try:
                self.assertEqual(workbook["ELM-RPE"]["B3"].value, 30)
                self.assertAlmostEqual(workbook["ONL"]["B3"].value, (300 + 400 + 700) / 3)
                log = workbook["Extraction_log"]
                self.assertEqual(log["G3"].value, 3)
                self.assertEqual(log["G4"].value, 3)
            finally:
                workbook.close()
            self.assertEqual(
                sorted(path.name for path in side.iterdir()),
                ["_thickness_vs_distance_from_fovea_LIGHT.txt", "rr_MCPAR.xlsx"],
            )

    def test_compile_matches_legacy_layout_and_reads_current_step4_results(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            subject = _build_subject(root)
            output = root / DEFAULT_OUTPUT_FILENAME
            log_messages = []
            progress_events = []

            result = compile_step4_results(
                root,
                output,
                include_fovea=True,
                log_callback=log_messages.append,
                progress_callback=lambda completed, total, message: progress_events.append(
                    (completed, total, message)
                ),
            )

            self.assertEqual(result.output_path, output)
            self.assertEqual(result.le_ids, ("7_0_0",))
            self.assertEqual(result.re_ids, ())
            self.assertEqual(
                result.measures,
                (
                    MeasureCompilation("rrMCP-AR", 1, 0),
                    MeasureCompilation("ELM-RPE", 1, 0),
                    MeasureCompilation("ONL", 1, 0),
                ),
            )
            self.assertTrue(any("LE subject(s) found: 7_0_0" in line for line in log_messages))
            self.assertTrue(log_messages[-1].endswith(str(output)))
            self.assertEqual(progress_events[0], (0, 1, "Scanning input folders..."))
            self.assertEqual(progress_events[1][:2], (1, 5))
            self.assertEqual(
                [event[:2] for event in progress_events if event[2].startswith("Processed ")],
                [(2, 5), (3, 5), (4, 5)],
            )
            self.assertEqual(progress_events[-2], (4, 5, "Saving workbook..."))
            self.assertEqual(progress_events[-1], (5, 5, "Saved compiled workbook."))
            self.assertEqual(
                [completed for completed, _total, _message in progress_events],
                sorted(completed for completed, _total, _message in progress_events),
            )

            workbook = load_workbook(output, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["rrMCP-AR", "ELM-RPE", "ONL", "Extraction_log"],
                )

                rr_sheet = workbook["rrMCP-AR"]
                self.assertEqual(rr_sheet.max_row, 22)
                self.assertEqual(rr_sheet["A1"].value, "LE")
                self.assertIsNone(rr_sheet["E1"].value)
                self.assertEqual(
                    [rr_sheet.cell(2, column).value for column in range(1, 8)],
                    ["ID", "Nasal", "Temporal", None, None, None, None],
                )
                self.assertEqual(rr_sheet["A3"].value, "7_0_0")
                self.assertEqual(rr_sheet["B3"].value, 1)
                self.assertEqual(rr_sheet["B22"].value, 20)
                self.assertEqual(rr_sheet["C3"].value, 101)
                self.assertEqual(rr_sheet["C22"].value, 120)
                self.assertTrue(rr_sheet["A3"].font.bold)
                self.assertEqual(rr_sheet["A3"].alignment.horizontal, "center")
                self.assertEqual(rr_sheet["A3"].alignment.vertical, "center")
                self.assertEqual(rr_sheet["B3"].number_format, "0.000")
                for column in ("A", "B", "C"):
                    self.assertEqual(rr_sheet.column_dimensions[column].width, 12.0)

                elm_sheet = workbook["ELM-RPE"]
                self.assertEqual(elm_sheet.max_row, 26)
                self.assertEqual(elm_sheet["A3"].value, "7_0_0")
                self.assertAlmostEqual(elm_sheet["B3"].value, 55.77)
                self.assertAlmostEqual(elm_sheet["C3"].value, 66.88)
                self.assertEqual(elm_sheet["B4"].value, 10.0)
                self.assertEqual(elm_sheet["B26"].value, 10.0)
                self.assertEqual(elm_sheet["C4"].value, 30.0)

                onl_sheet = workbook["ONL"]
                self.assertEqual(onl_sheet.max_row, 25)
                self.assertEqual(onl_sheet["A3"].value, "7_0_0")
                self.assertEqual(onl_sheet["B3"].value, 20.0)
                self.assertEqual(onl_sheet["B25"].value, 20.0)
                self.assertEqual(onl_sheet["C3"].value, 40.0)

                log_sheet = workbook["Extraction_log"]
                self.assertEqual(log_sheet.max_row, 4)
                self.assertEqual(log_sheet.max_column, 9)
                self.assertEqual(log_sheet["A2"].value, "rrMCP-AR")
                self.assertEqual(log_sheet["B2"].value, "LE")
                self.assertEqual(log_sheet["C2"].value, "7_0_0")
                self.assertEqual(log_sheet["I2"].value, "OK")
                self.assertTrue(str(log_sheet["E2"].value).endswith("rr_MCPAR.xlsx"))
                self.assertEqual(log_sheet["G2"].value, 20)
                self.assertEqual(log_sheet["G3"].value, 2300)
                self.assertTrue(log_sheet["A1"].font.bold)
                self.assertEqual(log_sheet["A1"].alignment.horizontal, "center")
                for column in "ABCDEFGHI":
                    self.assertEqual(log_sheet.column_dimensions[column].width, 28.0)
            finally:
                workbook.close()

            self.assertFalse(
                (subject / "nasal" / "_thickness_vs_distance_from_fovea_LIGHT2.txt").is_file()
            )
            self.assertFalse(
                (subject / "temporal" / "_thickness_vs_distance_from_fovea_LIGHT2.txt").is_file()
            )

    def test_fovea_can_be_excluded_from_elmrpe(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            _build_subject(root)
            output = root / "without_fovea.xlsx"

            compile_step4_results(root, output, include_fovea=False)

            workbook = load_workbook(output, data_only=True)
            try:
                elm_sheet = workbook["ELM-RPE"]
                self.assertEqual(elm_sheet.max_row, 25)
                self.assertEqual(elm_sheet["B3"].value, 10.0)
                self.assertEqual(elm_sheet["B25"].value, 10.0)
            finally:
                workbook.close()

    def test_current_step4_workbook_takes_precedence_over_legacy_results(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            _write_results(folder / "Results.xlsx", 100)
            current = folder / "rr_MCPAR.xlsx"
            _write_results(current, 1)

            self.assertEqual(Path(find_results_file(folder)), current)

    def test_skipped_subject_still_advances_each_measurement_progress_unit(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            nasal = root / "LE" / "8_0_0" / "nasal"
            nasal.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.title = "in"
            workbook.active.cell(1, 7).value = "Round"
            workbook.save(nasal / "Results.xlsx")
            workbook.close()
            output = root / "skipped.xlsx"
            progress_events = []

            result = compile_step4_results(
                root,
                output,
                progress_callback=lambda completed, total, message: progress_events.append(
                    (completed, total, message)
                ),
            )

            self.assertEqual(
                result.measures,
                (
                    MeasureCompilation("rrMCP-AR", 0, 1),
                    MeasureCompilation("ELM-RPE", 0, 1),
                    MeasureCompilation("ONL", 0, 1),
                ),
            )
            self.assertEqual(
                [event[:2] for event in progress_events if event[2].startswith("Processed ")],
                [(2, 5), (3, 5), (4, 5)],
            )
            self.assertEqual(progress_events[-1][:2], (5, 5))

    def test_save_failure_does_not_report_full_progress(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            _build_subject(root)
            output = root / "locked.xlsx"
            progress_events = []

            with mock.patch(
                "aidas.services.step4_compiler.Workbook.save",
                side_effect=PermissionError("workbook is locked"),
            ):
                with self.assertRaisesRegex(PermissionError, "workbook is locked"):
                    compile_step4_results(
                        root,
                        output,
                        progress_callback=lambda completed, total, message: (
                            progress_events.append((completed, total, message))
                        ),
                    )

            self.assertEqual(progress_events[-1], (4, 5, "Saving workbook..."))
            self.assertFalse(
                any(completed == total for completed, total, _message in progress_events)
            )

    def test_empty_or_flat_folder_is_rejected_instead_of_saving_empty_output(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            _write_results(root / "Results.xlsx", 1)
            output = root / "empty.xlsx"

            with self.assertRaisesRegex(ValueError, "No subject folders"):
                compile_step4_results(root, output)

            self.assertFalse(output.exists())

    def test_recursive_groups_nested_workbooks_and_extra_subject_folders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            expected = {
                "Control": ("Subject_2", "Subject_10"),
                "Treatment": ("Subject_2",),
                "Validation": ("001_extra",),
            }
            sources = {}
            for group, ids in expected.items():
                for subject_id in ids:
                    subject = root / "Experiment" / "Visit 1" / group / subject_id
                    for side, value in (("Nasal", 1), ("TEMPORAL", 101)):
                        result_dir = subject / side / "exports" / "final"
                        result_dir.mkdir(parents=True)
                        source = result_dir / "measurements.xlsx"
                        _write_results(source, value)
                        sources[group, subject_id, side] = str(source)
                    images = subject / "Images" / "Previews"
                    images.mkdir(parents=True)
                    (images / "scan.png").write_bytes(b"preview")
                    notes = subject / "Notes"
                    notes.mkdir()
                    (notes / "notes.txt").write_text("notes", encoding="utf-8")

            # These branches must never create subjects or group blocks.
            thickness_only = root / "Other" / "TextOnly" / "nasal"
            thickness_only.mkdir(parents=True)
            _write_thickness(thickness_only / "_thickness_vs_distance_from_fovea_LIGHT.txt", 8, 9)
            unrelated = root / "Other" / "Unrelated" / "Documents"
            unrelated.mkdir(parents=True)
            _write_results(unrelated / "notes.xlsx", 900)
            locked = root / "Other" / "LockOnly" / "temporal"
            locked.mkdir(parents=True)
            (locked / "~$rr_MCPAR.xlsx").write_bytes(b"lock")
            output = root / DEFAULT_OUTPUT_FILENAME
            result = compile_step4_results(root, output)

            self.assertEqual(result.group_ids, expected)
            self.assertEqual(result.measures[0], MeasureCompilation("rrMCP-AR", 4, 0))
            with_workbook = load_workbook(output, data_only=True)
            try:
                sheet = with_workbook["rrMCP-AR"]
                for column, (group, ids) in zip((1, 5, 9), expected.items()):
                    self.assertEqual(sheet.cell(1, column).value, group)
                    self.assertEqual(sheet.cell(2, column).value, "ID")
                    for index, subject_id in enumerate(ids):
                        row = 3 + index * 20
                        self.assertEqual(sheet.cell(row, column).value, subject_id)
                        self.assertEqual(sheet.cell(row, column + 1).value, 1)
                        self.assertEqual(sheet.cell(row, column + 2).value, 101)
                log = with_workbook["Extraction_log"]
                self.assertEqual(log["B1"].value, "Group")
                rr_rows = [row for row in log.iter_rows(min_row=2, values_only=True) if row[0] == "rrMCP-AR"]
                self.assertEqual(len(rr_rows), 4)
                for row in rr_rows:
                    self.assertEqual(row[4], sources[row[1], row[2], "Nasal"])
                    self.assertEqual(row[5], sources[row[1], row[2], "TEMPORAL"])
            finally:
                with_workbook.close()

            # Re-running must not discover the combined output as an input.
            self.assertEqual(compile_step4_results(root, output).group_ids, expected)

    def test_duplicate_group_and_subject_names_are_reported_without_losing_data(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            for batch in ("Batch1", "Batch2"):
                folder = root / batch / "Control" / "Subject_A" / "nasal"
                folder.mkdir(parents=True)
                _write_results(folder / "rr_MCPAR.xlsx", 1)
            output = root / "combined.xlsx"
            with self.assertRaisesRegex(ValueError, "Duplicate subject 'Subject_A' in group 'Control'"):
                compile_step4_results(root, output)
            self.assertFalse(output.exists())

    def test_temporal_only_subject_keeps_nasal_measurements_blank(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            side = root / "Temporal group" / "Subject_A" / "temporal"
            side.mkdir(parents=True)
            _write_results(side / "rr_MCPAR.xlsx", 101)
            _write_thickness(side / "_thickness_vs_distance_from_fovea_LIGHT.txt", 30, 40)
            output = root / "compiled.xlsx"
            result = compile_step4_results(root, output, include_fovea=False)
            self.assertEqual(result.group_ids, {"Temporal group": ("Subject_A",)})
            workbook = load_workbook(output, data_only=True)
            try:
                for measure, value in (("rrMCP-AR", 101), ("ELM-RPE", 30), ("ONL", 40)):
                    self.assertIsNone(workbook[measure]["B3"].value)
                    self.assertEqual(workbook[measure]["C3"].value, value)
            finally:
                workbook.close()

    def test_full_names_that_start_with_equals_are_written_as_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            side = root / "=Control" / "=Subject_A" / "nasal"
            side.mkdir(parents=True)
            _write_results(side / "rr_MCPAR.xlsx", 1)
            output = root / "compiled.xlsx"
            compile_step4_results(root, output)
            workbook = load_workbook(output, data_only=False)
            try:
                for cell, value in (
                    (workbook["rrMCP-AR"]["A1"], "=Control"),
                    (workbook["rrMCP-AR"]["A3"], "=Subject_A"),
                    (workbook["Extraction_log"]["B2"], "=Control"),
                    (workbook["Extraction_log"]["C2"], "=Subject_A"),
                ):
                    self.assertEqual(cell.value, value)
                    self.assertEqual(cell.data_type, "s")
            finally:
                workbook.close()

    def test_nested_current_workbook_is_preferred_to_direct_legacy_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            _write_results(folder / "Results.xlsx", 100)
            nested = folder / "exports"
            nested.mkdir()
            current = nested / "rr_MCPAR.xlsx"
            _write_results(current, 1)
            self.assertEqual(Path(find_results_file(folder)), current)

    def test_output_cannot_replace_an_input_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = _build_subject(root) / "nasal" / "rr_MCPAR.xlsx"
            original = source.read_bytes()
            with self.assertRaisesRegex(ValueError, "different from the input"):
                compile_step4_results(root, source)
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
