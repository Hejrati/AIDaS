from __future__ import annotations

import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

from aidas.services.measurement_inventory import InventoryCancelled, scan_measurement_directory
from aidas.services.step4_compiler import (
    compile_step4_results,
    find_light_files,
    find_results_file,
    find_subject_folders,
    find_thickness_file,
)
from aidas.utils.filesystem import walk_unique_directories


class Step5ScanCycleTests(unittest.TestCase):
    def test_mocked_parent_junction_is_pruned_before_it_can_recurse(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            thickness = root / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            thickness.touch()
            workbook = root / "rr_MCPAR.xlsx"
            workbook.touch()
            original_realpath = os.path.realpath

            def realpath(path, *args, **kwargs):
                if Path(path).name == "back":
                    return str(root)
                return original_realpath(path, *args, **kwargs)

            def walk(path, *, onerror=None, **_options):
                pending = [os.fspath(path)]
                visits = 0
                while pending:
                    directory = pending.pop()
                    visits += 1
                    self.assertLessEqual(visits, 3, "A parent junction was followed repeatedly")
                    subdirectories = ["back"]
                    yield directory, subdirectories, [workbook.name, thickness.name]
                    pending.extend(os.path.join(directory, name) for name in subdirectories)

            with mock.patch("aidas.utils.filesystem.os.path.realpath", side_effect=realpath), \
                    mock.patch("aidas.utils.filesystem.os.walk", side_effect=walk):
                rows = list(walk_unique_directories(root))
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0][1], [])
                self.assertEqual(Path(find_results_file(root)), workbook)
                self.assertEqual(Path(find_thickness_file(root)), thickness)
                inventory = scan_measurement_directory(root)
                self.assertEqual((inventory.workbooks, inventory.thickness_files), (1, 1))
                self.assertFalse(any(child.kind == "folder" for child in inventory.root.children))

    def test_mocked_duplicate_directory_alias_is_counted_once(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            physical = root / "physical"
            physical.mkdir()
            original_realpath = os.path.realpath

            def realpath(path, *args, **kwargs):
                if Path(path).name in ("first", "second"):
                    return str(physical)
                return original_realpath(path, *args, **kwargs)

            def walk(path, *, onerror=None, **_options):
                branches = ["first", "second"]
                yield str(root), branches, []
                for branch in branches:
                    yield str(root / branch), [], ["rr_MCPAR.xlsx"]

            with mock.patch("aidas.utils.filesystem.os.path.realpath", side_effect=realpath), \
                    mock.patch("aidas.utils.filesystem.os.walk", side_effect=walk):
                rows = list(walk_unique_directories(root))
                self.assertEqual([Path(row[0]).name for row in rows], [root.name, "first"])
                self.assertEqual(scan_measurement_directory(root).workbooks, 1)

    def test_file_ids_also_detect_aliases_with_different_canonical_names(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            physical = root / "physical"
            physical.mkdir()
            original_stat = os.stat
            physical_status = physical.stat()

            def stat(path, *args, **kwargs):
                if Path(path).name in ("first", "second"):
                    return physical_status
                return original_stat(path, *args, **kwargs)

            def walk(path, *, onerror=None, **_options):
                branches = ["first", "second"]
                yield str(root), branches, []
                for branch in branches:
                    yield str(root / branch), [], ["rr_MCPAR.xlsx"]

            with mock.patch("aidas.utils.filesystem.os.stat", side_effect=stat), \
                    mock.patch("aidas.utils.filesystem.os.walk", side_effect=walk):
                rows = list(walk_unique_directories(root))
            self.assertEqual([Path(row[0]).name for row in rows], [root.name, "first"])

    def test_walk_respects_caller_sorting_and_pruning(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            for name in ("skip", "B", "A"):
                (root / name / "nested").mkdir(parents=True)
            visited = []
            for directory, branches, _files in walk_unique_directories(root):
                visited.append(Path(directory).relative_to(root).as_posix())
                if Path(directory) == root:
                    branches[:] = sorted(name for name in branches if name != "skip")
                else:
                    branches.clear()
            self.assertEqual(visited, [".", "A", "B"])

    def test_real_link_cycle_does_not_repeat_files(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            side = root / "Control" / "Subject_A" / "nasal"
            side.mkdir(parents=True)
            (side / "rr_MCPAR.xlsx").touch()
            thickness = side / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            thickness.touch()
            link = side / "back"
            junction = False
            try:
                link.symlink_to(side, target_is_directory=True)
            except (OSError, NotImplementedError) as exc:
                if os.name != "nt":
                    self.skipTest(f"Directory symlinks are unavailable: {exc}")
                # Windows junctions need no administrator/symlink privilege.
                created = subprocess.run(
                    ["cmd", "/c", "mklink", "/J", str(link), str(side)],
                    capture_output=True, text=True, check=False,
                    creationflags=subprocess.CREATE_NO_WINDOW,
                )
                if created.returncode:
                    self.skipTest(f"Directory links are unavailable: {created.stderr.strip()}")
                junction = True
            try:
                rows = list(walk_unique_directories(root, followlinks=True))
                self.assertEqual(len(rows), 4)
                self.assertEqual(Path(find_thickness_file(root)), thickness)
                self.assertEqual(find_subject_folders(root), {"Control": {"Subject_A": str(side.parent)}})
                inventory = scan_measurement_directory(root)
                self.assertEqual((inventory.workbooks, inventory.thickness_files), (1, 1))
            finally:
                if junction:
                    link.rmdir()
                else:
                    link.unlink()

    def test_results_deeper_than_five_folders_still_compile_and_display(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            container = root.joinpath(*(f"level{index}" for index in range(7)))
            subject = container / "Control" / "Subject_A"
            side = subject / "nasal"
            exports = side.joinpath(*(f"export{index}" for index in range(7)))
            exports.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.cell(1, 7).value = "Round"
            workbook.active.cell(2, 7).value = 7.25
            workbook.save(exports / "rr_MCPAR.xlsx")
            workbook.close()
            thickness = exports / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            thickness.write_text("Distance\tWhole\tELM\tONL\n0\t0\t30\t40\n", encoding="utf-8")
            output = root / "compiled.xlsx"

            result = compile_step4_results(root, output, include_fovea=False)

            self.assertEqual(result.group_ids, {"Control": ("Subject_A",)})
            compiled = load_workbook(output, data_only=True)
            try:
                for sheet, value in (("rrMCP-AR", 7.25), ("ELM-RPE", 30), ("ONL", 40)):
                    self.assertEqual(compiled[sheet]["B3"].value, value)
            finally:
                compiled.close()
            self.assertEqual(find_light_files(subject), (str(thickness), None))
            inventory = scan_measurement_directory(root, output_path=output)
            self.assertEqual((inventory.workbooks, inventory.thickness_files), (1, 1))

    def test_thickness_search_preserves_priorities_and_ignores_hidden_profiles(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            legacy = root / "_thickness_vs_distance_from_fovea_LIGHT2.txt"
            legacy.touch()
            (root / "A_LIGHT.txt").touch()
            (root / "_fovea_light_profiles_LIGHT.txt").touch()
            hidden = root / ".hidden"
            hidden.mkdir()
            (hidden / "_thickness_vs_distance_from_fovea_LIGHT.txt").touch()
            self.assertEqual(Path(find_thickness_file(root)), legacy)
            current = root / "exports" / "_thickness_vs_distance_from_fovea_LIGHT.txt"
            current.parent.mkdir()
            current.touch()
            self.assertEqual(Path(find_thickness_file(root)), current)

    def test_inventory_cancels_while_scanning_nested_results(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "Control" / "Subject_A" / "nasal").mkdir(parents=True)
            calls = 0

            def cancelled():
                nonlocal calls
                calls += 1
                return calls == 2

            with self.assertRaises(InventoryCancelled):
                scan_measurement_directory(root, cancelled=cancelled)
            self.assertEqual(calls, 2)


if __name__ == "__main__":
    unittest.main()
